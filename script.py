from openpyxl import load_workbook

wb = load_workbook('input.xlsx')
price_st = wb.active

wb2 = load_workbook('output.xlsx')
price_our = wb2.active


def fill_price():
    for article in price_our['E5':'E4088']:
        art = str(article[0].value).strip()
        for article2 in price_st['B4':'B4687']:
            art2 = str(article2[0].value).strip()
            if art == art2:
                price_our['G' + str(article[0].row)].value = str(price_st['B' + str(article2[0].row)].value).strip()
                price_our['H' + str(article[0].row)].value = str(price_st['C' + str(article2[0].row)].value).strip()
                price_our['I' + str(article[0].row)].value = str(price_st['D' + str(article2[0].row)].value).strip()
                tmp = str(price_st['E' + str(article2[0].row)].value).strip()
                price_our['J' + str(article[0].row)].value = '' if tmp == 'None' else tmp
                break

    wb2.save('result.xlsx')


def find_diff():
    with open('diff.txt', 'w+') as diff:
        for article in price_st['B4':'B4651']:
            art_st = str(article[0].value).strip()
            found = False
            for article2 in price_our['E5':'E3947']:
                art_our = str(article2[0].value).strip()
                if art_st == art_our:
                    found = True
            if not found:
                diff.write(art_st + '\n')


fill_price()
find_diff()
